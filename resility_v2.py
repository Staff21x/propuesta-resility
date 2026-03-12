"""
Resility V2 — Correos ultra-personalizados con dato concreto de la web
Flujo: scrape web → extraer dato concreto → Claude redacta correo con ese dato en línea 1
"""

import os
import re
import time
import requests
from bs4 import BeautifulSoup
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import anthropic

# ─────────────────────────────────────────────
#  CONTACTOS
# ─────────────────────────────────────────────
CONTACTS = [
    ("Berenguer Mallol",  "Director General",   "MCA Chile Corredores de Seguros",      "bmallol@mcaseguros.cl",                "+56 9 8394885",  "Seguros"),
    ("Roberto Gatica",    "Gerente",             "Cono Sur Corredores de Seguros",       "roberto.gatica@conosurseguros.cl",     "+56 9 1294214",  "Seguros"),
    ("Cesar Acevedo",     "Gerente General",     "Phersu Corredores de Seguros",         "cesar.acevedo@phersu.cl",              "+56 9 99977211", "Seguros"),
    ("Oscar Harder",      "Gerente",             "Seguros Harder",                       "oficina@segurosharder.cl",             "+56 9 9161735",  "Seguros"),
    ("Isidoro Parraguez", "Director Ejecutivo",  "CP Brokers Corredores de Seguros",     "isidoro.parraguez@cpbrokers.cl",       "+56 9 9180543",  "Seguros"),
]

RESILITY_CONTEXT = """
Resility es un CyberSOC (Centro de Operaciones de Ciberseguridad) con Inteligencia Artificial en Chile:
- Monitoreo 24/7 con IA propia entrenada en amenazas latinoamericanas
- MDR (Managed Detection & Response): detecta y contiene incidentes en < 15 minutos
- Cumplimiento Ley 21.663 (Chile) y CMF sin armar equipo propio
- Integración con AWS, Azure, GCP
- Analistas certificados CISSP, CEH, OSCP
- Costo de SOC externo vs equipo interno: 80% más barato
"""

HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/120.0.0.0 Safari/537.36"
    ),
    "Accept-Language": "es-CL,es;q=0.9,en;q=0.8",
}

PROHIBITED_WORDS = ["solución", "regulatoria", "satisfacción", "optimizar", "potenciar", "cumplimiento"]


# ─────────────────────────────────────────────
#  SCRAPING
# ─────────────────────────────────────────────

SCRAPE_HEADERS_VARIANTS = [
    # Chrome normal
    {
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36",
        "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
        "Accept-Language": "es-CL,es;q=0.9",
        "Accept-Encoding": "gzip, deflate, br",
        "Connection": "keep-alive",
        "Upgrade-Insecure-Requests": "1",
    },
    # Firefox
    {
        "User-Agent": "Mozilla/5.0 (X11; Linux x86_64; rv:121.0) Gecko/20100101 Firefox/121.0",
        "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,image/avif,image/webp,*/*;q=0.8",
        "Accept-Language": "es-CL,es;q=0.8,en-US;q=0.5,en;q=0.3",
        "Connection": "keep-alive",
    },
    # Googlebot
    {
        "User-Agent": "Mozilla/5.0 (compatible; Googlebot/2.1; +http://www.google.com/bot.html)",
        "Accept": "text/html",
    },
]

EXTRA_PATHS = ["/quienes-somos", "/nosotros", "/sobre-nosotros", "/servicios", "/about"]


def scrape_website(domain: str) -> tuple[str, str]:
    """Scrapea la home (y rutas alternativas) probando distintos User-Agents."""
    base_urls = [f"https://www.{domain}", f"https://{domain}"]

    for base in base_urls:
        urls_to_try = [base] + [base.rstrip("/") + p for p in EXTRA_PATHS]
        for url in urls_to_try:
            for hdrs in SCRAPE_HEADERS_VARIANTS:
                try:
                    resp = requests.get(url, headers=hdrs, timeout=12,
                                        allow_redirects=True, verify=True)
                    if resp.status_code == 200 and len(resp.text) > 200:
                        soup = BeautifulSoup(resp.text, "lxml")
                        for tag in soup(["script", "style", "nav", "footer", "header", "aside"]):
                            tag.decompose()
                        text = soup.get_text(separator=" ", strip=True)
                        text = re.sub(r"\s+", " ", text)
                        if len(text) > 100:
                            return url, text[:3000]
                except Exception:
                    continue
            time.sleep(0.5)

    return f"https://www.{domain}", ""


def extract_concrete_fact(nombre: str, empresa: str, web_url: str, web_text: str, client: anthropic.Anthropic) -> str:
    """Pide a Claude que extraiga UN dato concreto y verificable del texto web."""
    if not web_text.strip():
        return ""

    prompt = f"""Del siguiente texto extraído de la web de "{empresa}", extrae UN SOLO dato concreto y específico.

Prioridad (en orden):
1. Un servicio o especialidad concreta que ofrecen (ej: "seguros de vida colectivos para pymes")
2. Una frase real de su home que describa qué hacen (máx 15 palabras, entre comillas)
3. Años en el mercado o año de fundación
4. Un nicho o tipo de cliente específico que mencionan
5. Una cobertura o producto distintivo

REGLAS:
- Devuelve SOLO el dato, sin explicación ni contexto adicional
- Máximo 20 palabras
- Si no hay nada concreto y verificable, responde exactamente: SIN_DATO

Texto web:
{web_text[:2000]}"""

    msg = client.messages.create(
        model="claude-opus-4-6",
        max_tokens=100,
        messages=[{"role": "user", "content": prompt}],
    )
    dato = msg.content[0].text.strip()
    return "" if dato == "SIN_DATO" else dato


# ─────────────────────────────────────────────
#  LINKEDIN VERIFICATION
# ─────────────────────────────────────────────

def verify_linkedin(nombre: str, empresa: str) -> str:
    """Busca en DuckDuckGo si la persona aparece activa en LinkedIn (2025/2026)."""
    query = f'"{nombre}" "{empresa}" linkedin 2025 OR 2026'
    url = f"https://html.duckduckgo.com/html/?q={requests.utils.quote(query)}"
    try:
        resp = requests.get(url, headers=HEADERS, timeout=12)
        soup = BeautifulSoup(resp.text, "lxml")
        titles   = [t.get_text(strip=True).lower() for t in soup.select(".result__title")[:5]]
        snippets = [s.get_text(strip=True).lower() for s in soup.select(".result__snippet")[:5]]
        all_text = " ".join(titles + snippets)
        hrefs    = " ".join(a.get("href", "") for a in soup.select(".result__title a"))

        nombre_partes  = nombre.lower().split()
        empresa_lower  = empresa.lower().split()[0]
        menciona_persona  = any(p in all_text for p in nombre_partes if len(p) > 3)
        menciona_empresa  = empresa_lower in all_text
        menciona_reciente = "2025" in all_text or "2026" in all_text
        es_linkedin       = "linkedin" in all_text or "linkedin.com" in hrefs

        if menciona_persona and (menciona_empresa or es_linkedin) and menciona_reciente:
            return "✅ Verificado"
        return "⚠️ Verificar manualmente"
    except Exception:
        return "⚠️ Verificar manualmente"


# ─────────────────────────────────────────────
#  GENERACIÓN DE CORREO
# ─────────────────────────────────────────────

def generate_email(contact: dict, web_url: str, dato_concreto: str, client: anthropic.Anthropic) -> tuple[str, str]:
    """Claude genera asunto + correo de 6 líneas máximo."""

    primera_linea_instruccion = (
        f'La PRIMERA línea del cuerpo DEBE mencionar este dato concreto de {contact["empresa"]}: "{dato_concreto}"'
        if dato_concreto
        else f'La PRIMERA línea del cuerpo debe hacer referencia al rubro específico de {contact["empresa"]} como corredora de seguros, mostrando que investigaste'
    )

    prohibidas = ", ".join(f'"{w}"' for w in PROHIBITED_WORDS)

    prompt = f"""Eres el Director Comercial de Resility, CyberSOC con IA, Chile.
Redacta un correo de ventas B2B para este prospecto.

PROSPECTO:
- Nombre: {contact['nombre']}
- Cargo: {contact['cargo']}
- Empresa: {contact['empresa']}
- Web: {web_url}

SOBRE RESILITY:
{RESILITY_CONTEXT}

CONTEXTO — CORREDORAS DE SEGUROS:
Manejan RUT, datos de salud, pólizas y cuentas bancarias de clientes.
Sin equipo de ciberseguridad propio. Sujetas a Ley 21.663 con multas de hasta 40.000 UTM.
Un ransomware las paraliza en horas y destruye su reputación.

INSTRUCCIONES CORREO:
1. {primera_linea_instruccion}
2. Cuerpo: MÁXIMO 6 LÍNEAS en total (incluyendo saludo y firma)
3. Saludo: solo el primer nombre, sin "Estimado/a"
4. Una sola idea central: manejan datos críticos sin protección → Resility es su SOC externalizado
5. CTA: proponer llamada de 15 minutos esta semana
6. Firma: "Equipo Comercial · Resility | CyberSOC con IA"
7. Tono: colega ejecutivo, sin tecnicismos, sin puntos de bala

ASUNTO — REGLAS:
- MÁXIMO 7 PALABRAS. Sin excepción.
- Provocador, genera urgencia real, hace que el receptor piense "esto me afecta"
- PROHIBIDO usar estas palabras: {prohibidas}
- PROHIBIDO signos de exclamación
- PROHIBIDO empezar con el nombre de la empresa
- Diferente y único para este contacto específico

Devuelve EXACTAMENTE este formato (sin texto adicional):
ASUNTO: [asunto]
---
[cuerpo del correo]"""

    msg = client.messages.create(
        model="claude-opus-4-6",
        max_tokens=512,
        messages=[{"role": "user", "content": prompt}],
    )

    raw = msg.content[0].text.strip()
    if "---" in raw:
        parts = raw.split("---", 1)
        asunto = parts[0].replace("ASUNTO:", "").strip()
        cuerpo = parts[1].strip()
    else:
        lines = raw.split("\n")
        asunto = lines[0].replace("ASUNTO:", "").strip()
        cuerpo = "\n".join(lines[1:]).strip()

    return asunto, cuerpo


# ─────────────────────────────────────────────
#  EXCEL
# ─────────────────────────────────────────────

def build_excel(results: list, output_path: str):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Resility V2 — Corredoras"

    dark_fill   = PatternFill("solid", fgColor="0A1628")
    blue_fill   = PatternFill("solid", fgColor="1A56DB")
    alt_fill    = PatternFill("solid", fgColor="EEF2FF")
    green_fill  = PatternFill("solid", fgColor="D1FAE5")
    yellow_fill = PatternFill("solid", fgColor="FEF3C7")
    white_fill  = PatternFill("solid", fgColor="FFFFFF")

    hdr_font    = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    title_font  = Font(name="Calibri", bold=True, color="FFFFFF", size=13)
    normal_font = Font(name="Calibri", size=10)
    green_font  = Font(name="Calibri", size=10, color="065F46")
    yellow_font = Font(name="Calibri", size=10, color="92400E")
    wrap_align  = Alignment(wrap_text=True, vertical="top")
    center      = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style="thin",   color="D1D5DB"),
        right=Side(style="thin",  color="D1D5DB"),
        top=Side(style="thin",    color="D1D5DB"),
        bottom=Side(style="thin", color="D1D5DB"),
    )

    # Título
    ws.merge_cells("A1:M1")
    ws["A1"] = f"RESILITY V2 — Correos Personalizados con Dato Concreto  |  {datetime.now().strftime('%d/%m/%Y %H:%M')}"
    ws["A1"].font = title_font
    ws["A1"].fill = dark_fill
    ws["A1"].alignment = center

    columns = [
        "#", "Nombre", "Cargo", "Empresa", "Email", "Teléfono",
        "Web", "Dato Concreto Extraído", "Asunto", "Correo (6 líneas)",
        "Estado", "LinkedIn",
    ]
    widths = [4, 20, 20, 28, 30, 16, 28, 40, 38, 75, 10, 22]

    for ci, (col, w) in enumerate(zip(columns, widths), start=1):
        cell = ws.cell(row=2, column=ci, value=col)
        cell.font = hdr_font
        cell.fill = blue_fill
        cell.alignment = center
        cell.border = thin_border
        ws.column_dimensions[openpyxl.utils.get_column_letter(ci)].width = w

    ws.row_dimensions[1].height = 28
    ws.row_dimensions[2].height = 22

    for i, row in enumerate(results, start=1):
        r = i + 2
        base_fill = alt_fill if i % 2 == 0 else white_fill
        values = [
            i,
            row["nombre"], row["cargo"], row["empresa"],
            row["email"],  row["telefono"], row["web_url"],
            row["dato_concreto"],
            row["asunto"], row["correo"],
            row["estado"], row["linkedin"],
        ]
        for ci, val in enumerate(values, start=1):
            cell = ws.cell(row=r, column=ci, value=val)
            cell.border = thin_border

            # Columna 10: correo (wrap + alto)
            if ci == 10:
                cell.font = normal_font
                cell.fill = base_fill
                cell.alignment = wrap_align
                ws.row_dimensions[r].height = 110
            # Columna 8: dato concreto (wrap)
            elif ci == 8:
                cell.font = normal_font
                cell.fill = base_fill
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            # Columna 9: asunto (wrap)
            elif ci == 9:
                cell.font = normal_font
                cell.fill = base_fill
                cell.alignment = Alignment(wrap_text=True, vertical="center")
            # Columna 12: LinkedIn (color)
            elif ci == 12:
                cell.alignment = Alignment(horizontal="center", vertical="center")
                if str(val).startswith("✅"):
                    cell.fill = green_fill
                    cell.font = green_font
                else:
                    cell.fill = yellow_fill
                    cell.font = yellow_font
            else:
                cell.font = normal_font
                cell.fill = base_fill
                cell.alignment = Alignment(vertical="center")

    ws.freeze_panes = "A3"
    wb.save(output_path)
    print(f"\n✅  Excel guardado: {output_path}")


# ─────────────────────────────────────────────
#  MAIN
# ─────────────────────────────────────────────

def main():
    api_key = os.environ.get("ANTHROPIC_API_KEY", "")
    if not api_key:
        raise SystemExit("❌  ANTHROPIC_API_KEY no encontrada.\n   export ANTHROPIC_API_KEY='sk-ant-...'")

    client = anthropic.Anthropic(api_key=api_key)

    print("=" * 65)
    print("  RESILITY V2 — Corredoras de Seguros Chile")
    print(f"  {len(CONTACTS)} contactos  |  {datetime.now().strftime('%d/%m/%Y %H:%M')}")
    print("=" * 65)

    results = []

    for idx, (nombre, cargo, empresa, email, telefono, rubro) in enumerate(CONTACTS, start=1):
        contact = {"nombre": nombre, "cargo": cargo, "empresa": empresa,
                   "email": email, "telefono": telefono, "rubro": rubro}

        print(f"\n[{idx:02d}/{len(CONTACTS)}] {nombre} — {empresa}")
        print(f"         {cargo}")

        domain = email.split("@")[1]
        print(f"  🌐  Scrapeando {domain}...", end="", flush=True)
        web_url, web_text = scrape_website(domain)
        print(f" {len(web_text)} chars")

        print(f"  🔎  Extrayendo dato concreto...", end="", flush=True)
        dato_concreto = extract_concrete_fact(nombre, empresa, web_url, web_text, client) if web_text else ""
        print(f" {'«' + dato_concreto[:60] + '»' if dato_concreto else 'sin dato (web vacía)'}")

        print(f"  🔍  Verificando LinkedIn...", end="", flush=True)
        linkedin = verify_linkedin(nombre, empresa)
        print(f" {linkedin}")

        print(f"  🤖  Generando correo...", end="", flush=True)
        try:
            asunto, correo = generate_email(contact, web_url, dato_concreto, client)
            estado = "OK"
            print(" ✓")
        except Exception as e:
            asunto = "Error"
            correo = str(e)
            estado = "ERROR"
            print(f" ✗ {e}")

        print(f"\n  📧  {asunto}")
        print(f"  ✉️   {correo[:100].replace(chr(10), ' ')}...")

        results.append({
            **contact,
            "web_url":       web_url,
            "dato_concreto": dato_concreto or "(web sin contenido útil)",
            "asunto":        asunto,
            "correo":        correo,
            "estado":        estado,
            "linkedin":      linkedin,
        })

        if idx < len(CONTACTS):
            time.sleep(2)

    output_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "resility_v2_correos.xlsx")
    build_excel(results, output_path)

    print("\n" + "=" * 65)
    print(f"  COMPLETADO — {len(results)} correos  |  resility_v2_correos.xlsx")
    print("=" * 65)


if __name__ == "__main__":
    main()
